"""
generate_reports.py
===================
Converts research_paper.md and executive_summary.md to .docx format
using python-docx.

Usage:
    python generate_reports.py
"""

import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPORTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "reports")


def md_to_docx(md_path: str, docx_path: str, title: str = None):
    """
    Simple markdown-to-docx converter.
    Handles: headers (H1-H4), bold, italic, tables, bullet lists, horizontal rules.
    """
    doc = Document()

    # Style the default paragraph
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ("---", "===", "***"):
            doc.add_paragraph("─" * 50)
            i += 1
            continue

        # H1
        if line.startswith("# ") and not line.startswith("## "):
            p = doc.add_heading(line[2:].strip(), level=1)
            p.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
            i += 1
            continue

        # H2
        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        # H3
        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        # H4
        if line.startswith("#### "):
            p = doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        # Bullet list
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            text = line.strip()[2:]
            text = _clean_markdown(text)
            p = doc.add_paragraph(style="List Bullet")
            _add_run_with_formatting(p, text)
            i += 1
            continue

        # Numbered list
        if re.match(r"^\d+\.", line.strip()):
            text = re.sub(r"^\d+\.\s*", "", line.strip())
            text = _clean_markdown(text)
            p = doc.add_paragraph(style="List Number")
            _add_run_with_formatting(p, text)
            i += 1
            continue

        # Table (starts with |)
        if line.strip().startswith("|"):
            # Collect all table lines
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].rstrip())
                i += 1
            _add_table(doc, table_lines)
            continue

        # Bold/italic metadata lines (e.g., **Authors:** ...)
        if line.startswith("**") or line.startswith("*"):
            p = doc.add_paragraph()
            _add_run_with_formatting(p, _clean_markdown(line))
            i += 1
            continue

        # Normal paragraph
        text = _clean_markdown(line)
        if text.strip():
            p = doc.add_paragraph()
            _add_run_with_formatting(p, text)
        i += 1

    doc.save(docx_path)
    print(f"  Saved: {docx_path}")


def _clean_markdown(text: str) -> str:
    """Remove markdown link syntax, keeping text."""
    # Convert [text](url) → text
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text


def _add_run_with_formatting(paragraph, text: str):
    """Add a run with bold and italic formatting from markdown.**bold** *italic*"""
    # Split on **..** and *..* patterns
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def _add_table(doc, table_lines: list):
    """Add a docx table from markdown table lines."""
    rows = []
    for line in table_lines:
        # Skip separator rows (|---|---|)
        if re.match(r"^\|[-| :]+\|$", line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)

    if not rows:
        return

    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        for c_idx, cell_text in enumerate(row_data):
            if c_idx < n_cols:
                cell = row.cells[c_idx]
                cell.text = _clean_markdown(cell_text)
                if r_idx == 0:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True


def main():
    print("\nGenerating .docx reports...")

    md_to_docx(
        os.path.join(REPORTS_DIR, "research_paper.md"),
        os.path.join(REPORTS_DIR, "research_paper.docx"),
        title="Predictive Modeling and Risk Scoring for Bank Customer Churn",
    )

    md_to_docx(
        os.path.join(REPORTS_DIR, "executive_summary.md"),
        os.path.join(REPORTS_DIR, "executive_summary.docx"),
        title="Executive Summary — Customer Churn Intelligence Platform",
    )

    print("\n[OK] Reports generated!")


if __name__ == "__main__":
    main()
